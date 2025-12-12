from __future__ import annotations

from io import BytesIO

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment
from openpyxl.utils import get_column_letter


def generate_codtafsiltamin_template() -> BytesIO:
    """Generate Excel template for codtafsiltamin file."""
    headers = [
        "کد تامین کننده",
        "نام تامین کننده",
        "وضعیت",
        "نوع",
        "کد تفصیلی",
    ]
    
    sample_data = [
        ["732", "سیاوش ایوبی (تعمیر)", "فعال", "داخلی", "0026968"],
        ["731", "احمد رحیمیان (نمایندگی)", "فعال", "داخلی", "0026956"],
    ]
    
    return _create_excel_file(headers, sample_data, "codtafsiltamin")


def generate_contractors_one_template() -> BytesIO:
    """Generate Excel template for contractors-1 file."""
    # All columns from actual Excel file (excluding Unnamed: 0, کد تفصیلی, کد تامین کننده)
    headers = [
        "شماره اتوماسیون",
        "تاريخ ایجاد سند حسابداری فاکتور",
        "در اختیار ناظر",
        "تاریخ تحویل به حسابداری",
        "تاریخ تحویل به ناظر",
        "تاريخ ایجاد فاکتور",
        "شماره روکش",
        "بهره بردار",
        "موضوع هزینه فاکتور خرید",
        "مبلغ کل",
        "مبلغ بدون مالیات",
        "محدوده زمانی انجام کار",
        "وضعیت فاکتور",
        "توضیحات فاکتور خرید",
        "شماره صورت حساب پیمانکار فاکتور خرید",
        "شماره مجوز فاکتور خرید",
        "نوع مجوز فاکتور خرید",
        "نام تامین کننده",
        "شماره قرارداد کارفرما فاکتور خرید",
        "نام کارفرما",
    ]
    
    sample_data = [
        [
            "",  # شماره اتوماسیون
            "1401/09/12",  # تاريخ ایجاد سند حسابداری فاکتور
            "",  # در اختیار ناظر
            "",  # تاریخ تحویل به حسابداری
            "",  # تاریخ تحویل به ناظر
            "",  # تاريخ ایجاد فاکتور
            "14010010",  # شماره روکش
            "حوزه مدیرعامل",  # بهره بردار
            "",  # موضوع هزینه فاکتور خرید
            "1084877000",  # مبلغ کل
            "995300000",  # مبلغ بدون مالیات
            "",  # محدوده زمانی انجام کار
            "تایید شده",  # وضعیت فاکتور
            "",  # توضیحات فاکتور خرید
            "",  # شماره صورت حساب پیمانکار فاکتور خرید
            "",  # شماره مجوز فاکتور خرید
            "",  # نوع مجوز فاکتور خرید
            "پیشرو فناوران فرتاک",  # نام تامین کننده
            "",  # شماره قرارداد کارفرما فاکتور خرید
            "",  # نام کارفرما
        ],
        [
            "",  # شماره اتوماسیون
            "1401/08/29",  # تاريخ ایجاد سند حسابداری فاکتور
            "",  # در اختیار ناظر
            "",  # تاریخ تحویل به حسابداری
            "",  # تاریخ تحویل به ناظر
            "",  # تاريخ ایجاد فاکتور
            "14010011",  # شماره روکش
            "حوزه مدیرعامل",  # بهره بردار
            "",  # موضوع هزینه فاکتور خرید
            "16490479200",  # مبلغ کل
            "15128880000",  # مبلغ بدون مالیات
            "",  # محدوده زمانی انجام کار
            "تایید شده",  # وضعیت فاکتور
            "",  # توضیحات فاکتور خرید
            "",  # شماره صورت حساب پیمانکار فاکتور خرید
            "",  # شماره مجوز فاکتور خرید
            "",  # نوع مجوز فاکتور خرید
            "واگن سازان تبریز",  # نام تامین کننده
            "",  # شماره قرارداد کارفرما فاکتور خرید
            "",  # نام کارفرما
        ],
    ]
    
    return _create_excel_file(headers, sample_data, "contractors-1")


def generate_contractors_two_template() -> BytesIO:
    """Generate Excel template for contractors-2 file."""
    headers = [
        "شماره",
        "تاریخ",
        "تامین کننده احد / رمز تامین",
        "وضعیت",
        "قلم خرید",
        "مبلغ ناخالص",
        "مبنا",
        "توضیحات",
    ]
    
    sample_data = [
        [
            "14040450",
            "1404/03/03",
            "محمد جواد م ریلی (محصول)",
            "ثبت شده",
            "لجستیک ریلی",
            "222000000",
            "14046576",
            "14040331",
        ],
        [
            "14040453",
            "1404/04/02",
            "سید جلال ح ابنیه و تاسیس",
            "ثبت شده",
            "ساخت سوله",
            "100000000",
            "14046575",
            "0452",
        ],
    ]
    
    return _create_excel_file(headers, sample_data, "contractors-2")


def _create_excel_file(headers: list[str], sample_data: list[list], sheet_name: str) -> BytesIO:
    """Create an Excel file with headers and sample data."""
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_name
    
    # Write headers
    for col_idx, header in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center")
    
    # Write sample data
    for row_idx, row_data in enumerate(sample_data, start=2):
        for col_idx, value in enumerate(row_data, start=1):
            ws.cell(row=row_idx, column=col_idx, value=value)
    
    # Adjust column widths
    for col_idx, header in enumerate(headers, start=1):
        col_letter = get_column_letter(col_idx)
        ws.column_dimensions[col_letter].width = max(len(str(header)) + 2, 15)
    
    # Save to BytesIO
    output = BytesIO()
    wb.save(output)
    output.seek(0)
    return output

